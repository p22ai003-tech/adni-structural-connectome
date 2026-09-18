#!/usr/bin/env python3
"""Build the figure-first results storyboard requested for supervisor review."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path("/home/ec2-user/exp/research_audit")
OUT = ROOT / "outputs" / "supervisor_revision_v4"
OUT.mkdir(parents=True, exist_ok=True)
PPTX = OUT / "structural_connectome_key_results_storyboard_v1.pptx"

FIG_DESIGN = ROOT / "figures" / "author_ready_study_design.png"
FIG_CONTEXT = ROOT / "figures" / "whole_brain_microstructure_forest_v2.png"
FIG_PRIMARY = (
    ROOT
    / "outputs"
    / "stage_specific_diffusivity_v2"
    / "figures"
    / "transition_signature_forest.png"
)
FIG_CENTRAL = ROOT / "figures" / "multivariate_stage_vector_v1.png"
FIG_SENSITIVITY = ROOT / "figures" / "author_ready_sensitivity_forest.png"

NAVY = RGBColor(27, 44, 70)
BLUE = RGBColor(38, 111, 165)
MAGENTA = RGBColor(185, 49, 113)
TEAL = RGBColor(25, 132, 141)
GREEN = RGBColor(44, 125, 92)
AMBER = RGBColor(184, 116, 25)
INK = RGBColor(34, 42, 52)
MID = RGBColor(91, 104, 119)
LIGHT = RGBColor(239, 243, 247)
LINE = RGBColor(207, 216, 225)
WHITE = RGBColor(255, 255, 255)


def set_run(run, size, color=INK, bold=False, font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    set_run(run, size=size, color=color, bold=bold)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=20, color=INK,
                   bullet=False, spacing=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    for index, item in enumerate(lines):
        if isinstance(item, tuple):
            lead, body = item
        else:
            lead, body = "", item
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.level = 0
        paragraph.text = ""
        paragraph.space_after = Pt(spacing)
        if bullet:
            paragraph.text = "• "
            set_run(paragraph.runs[0], size=size, color=color)
        if lead:
            r1 = paragraph.add_run()
            r1.text = lead
            set_run(r1, size=size, color=color, bold=True)
        r2 = paragraph.add_run()
        r2.text = body
        set_run(r2, size=size, color=color)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.55, 0.28, 12.2, 0.48, size=26, color=NAVY, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.82),
                                  Inches(1.25), Inches(0.06))
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 1.95, 0.69, 10.8, 0.32, size=12, color=MID)


def add_footer(slide, number, note="Historical-output discovery analysis"):
    add_text(slide, note, 0.58, 7.14, 10.9, 0.18, size=8.5, color=MID)
    add_text(slide, str(number), 12.15, 7.10, 0.55, 0.22, size=9, color=MID,
             align=PP_ALIGN.RIGHT)


def add_chip(slide, label, x, y, w, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(0.34))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_run(r, 10, color=WHITE, bold=True)
    return shape


def add_panel(slide, x, y, w, h, fill=WHITE, border=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)
    return shape


def add_image_contain(slide, path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py),
                                    width=Inches(pw), height=Inches(ph))


def add_takeaway(slide, text, x=0.65, y=6.45, w=12.0, h=0.48, color=TEAL):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT
    panel.line.color.rgb = color
    panel.line.width = Pt(1.5)
    add_text(slide, text, x + 0.15, y + 0.08, w - 0.3, h - 0.12,
             size=13, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)


def build():
    for figure in [FIG_DESIGN, FIG_CONTEXT, FIG_PRIMARY, FIG_CENTRAL, FIG_SENSITIVITY]:
        if not figure.exists():
            raise FileNotFoundError(figure)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Title
    slide = prs.slides.add_slide(blank)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    banner.line.fill.background()
    add_text(slide, "Key results and manuscript narrative", 0.72, 1.40, 11.9, 0.72,
             size=34, color=NAVY, bold=True)
    add_text(slide, "Distinct axial and radial diffusivity patterns across adjacent stages of cognitive decline",
             0.76, 2.30, 11.4, 0.75, size=23, color=INK)
    add_text(slide, "Sabeesh Ethiraj  |  Dipanjan Roy\nSchool of Artificial Intelligence and Data Science, IIT Jodhpur",
             0.76, 3.30, 8.5, 0.78, size=16, color=MID)
    add_chip(slide, "513 participants", 0.78, 4.55, 1.65, BLUE)
    add_chip(slide, "51 centres", 2.60, 4.55, 1.25, TEAL)
    add_chip(slide, "17 systems", 4.02, 4.55, 1.35, AMBER)
    add_chip(slide, "Figure-first working outline", 5.55, 4.55, 2.35, MAGENTA)
    add_text(slide,
             "Purpose: agree on the figures and the scientific argument before rewriting the manuscript.",
             0.78, 5.35, 11.5, 0.54, size=18, color=INK, bold=True)
    add_footer(slide, 1, "Supervisor-review storyboard | 19 July 2026")

    # 2. Questions and design
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The paper is organized around three questions",
              "These are post-hoc analysis questions, not preregistered hypotheses")
    add_panel(slide, 0.62, 1.10, 3.92, 1.10, fill=RGBColor(242, 248, 252), border=BLUE)
    add_panel(slide, 4.71, 1.10, 3.92, 1.10, fill=RGBColor(244, 250, 248), border=GREEN)
    add_panel(slide, 8.80, 1.10, 3.92, 1.10, fill=RGBColor(252, 246, 249), border=MAGENTA)
    add_text(slide, "1  Stage", 0.82, 1.28, 1.4, 0.28, size=18, color=BLUE, bold=True)
    add_text(slide, "Do CN→MCI and MCI→AD show different tensor-component patterns?",
             0.82, 1.58, 3.45, 0.44, size=14, color=INK)
    add_text(slide, "2  Spatial expression", 4.91, 1.28, 2.3, 0.28, size=18, color=GREEN, bold=True)
    add_text(slide, "Are the effects localized to a named network or distributed across systems?",
             4.91, 1.58, 3.45, 0.44, size=14, color=INK)
    add_text(slide, "3  Multisite robustness", 9.00, 1.28, 2.7, 0.28, size=18, color=MAGENTA, bold=True)
    add_text(slide, "Do common-site, site-deletion and held-out-site analyses retain the signal?",
             9.00, 1.58, 3.45, 0.44, size=14, color=INK)
    add_image_contain(slide, FIG_DESIGN, 0.62, 2.42, 12.08, 3.78)
    add_takeaway(slide, "Narrative rule: each Results subsection answers one question and points to one figure.")
    add_footer(slide, 2)

    # 3. Conventional context
    slide = prs.slides.add_slide(blank)
    add_title(slide, "First establish the familiar diffusion pattern",
              "FA, MD, AxD and RD across the three clinical contrasts")
    add_image_contain(slide, FIG_CONTEXT, 0.46, 1.05, 9.30, 5.55)
    add_panel(slide, 9.88, 1.20, 2.84, 4.90, fill=LIGHT)
    add_text(slide, "What this figure does", 10.10, 1.43, 2.40, 0.36,
             size=18, color=NAVY, bold=True)
    add_rich_lines(slide, [
        ("Context. ", "The cohort shows the expected broad diffusion abnormality pattern."),
        ("Early contrast. ", "MD is higher in MCI than CN after 12-test Holm correction."),
        ("Later disease. ", "MD, AxD and RD are clearly higher in AD than CN."),
        ("Not the novelty. ", "This figure establishes face validity; it does not carry the central claim."),
    ], 10.10, 1.98, 2.40, 3.75, size=14, spacing=10)
    add_takeaway(slide, "The conventional metrics set the biological context before the adjacent-stage analysis.")
    add_footer(slide, 3)

    # 4. Primary finding
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Different tensor components mark the two adjacent-stage contrasts",
              "The two endpoints were tested together within each support scope")
    add_image_contain(slide, FIG_PRIMARY, 0.42, 1.04, 8.55, 5.46)
    add_panel(slide, 9.13, 1.16, 3.55, 4.98, fill=WHITE)
    add_chip(slide, "MCI − CN", 9.42, 1.45, 1.20, BLUE)
    add_text(slide, "AxD β=0.294\nwild-Holm p=0.0046", 9.42, 1.88, 2.78, 0.70,
             size=18, color=INK, bold=True)
    add_chip(slide, "AD − MCI", 9.42, 2.86, 1.20, MAGENTA)
    add_text(slide, "RD β=0.591\nwild-Holm p=0.0164", 9.42, 3.29, 2.78, 0.70,
             size=18, color=INK, bold=True)
    add_text(slide, "Pair-common sites retain both effects. The all-three-site analysis remains positive but is borderline after finite-cluster correction (p=0.061).",
             9.42, 4.36, 2.86, 1.26, size=13.5, color=MID)
    add_takeaway(slide, "Central observation: AxD separates MCI from CN, whereas RD separates AD from MCI.", color=MAGENTA)
    add_footer(slide, 4)

    # 5. Central multiscale figure
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The central figure links stage, systems and centre robustness",
              "Machine learning appears only as a held-out-site test of the late RD pattern")
    add_image_contain(slide, FIG_CENTRAL, 0.40, 1.02, 12.45, 5.70)
    add_footer(slide, 5)

    # 6. Sensitivity
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The late RD effect is more sensitive to sample restrictions",
              "Filled points pass the jointly corrected two-endpoint test; open points do not")
    add_image_contain(slide, FIG_SENSITIVITY, 0.48, 1.00, 9.50, 5.60)
    add_panel(slide, 10.04, 1.20, 2.68, 4.95, fill=LIGHT)
    add_text(slide, "Interpretation", 10.27, 1.46, 2.18, 0.35, size=18, color=NAVY, bold=True)
    add_rich_lines(slide, [
        ("AxD. ", "Stable across the main completeness thresholds."),
        ("RD. ", "Supported in the primary and ≥9-system analyses, but attenuated in smaller restricted subsets."),
        ("Meaning. ", "The RD result is not erased; its precision and generality depend more strongly on cohort composition."),
        ("Writing choice. ", "Report this once as a limitation, not as a list of failed experiments."),
    ], 10.25, 1.98, 2.22, 3.72, size=13.5, spacing=10)
    add_takeaway(slide, "Sensitivity analyses qualify the main result without replacing it.")
    add_footer(slide, 6)

    # 7. What the data support
    slide = prs.slides.add_slide(blank)
    add_title(slide, "What the present data support—and what they do not")
    add_panel(slide, 0.68, 1.18, 5.88, 5.45, fill=RGBColor(243, 250, 247), border=GREEN)
    add_panel(slide, 6.78, 1.18, 5.88, 5.45, fill=RGBColor(252, 246, 247), border=MAGENTA)
    add_text(slide, "Supported in the historical cohort", 0.98, 1.48, 5.25, 0.45,
             size=21, color=GREEN, bold=True)
    add_rich_lines(slide, [
        "Positive AxD MCI−CN and RD AD−MCI effects after joint finite-cluster correction.",
        "Broad system participation: 17/17 AxD and 13/17 RD systems pass endpoint-specific FDR.",
        "All 51 paired site deletions retain both positive effects and nominal significance.",
        "The system-level RD model improves over demographic/acquisition baseline at held-out sites.",
    ], 1.00, 2.03, 5.18, 3.92, size=16, bullet=True, spacing=14)
    add_text(slide, "Not established", 7.08, 1.48, 5.20, 0.45,
             size=21, color=MAGENTA, bold=True)
    add_rich_lines(slide, [
        "A limbic, DMN, hub, long-range or hemispheric epicentre.",
        "A cellular interpretation of AxD or RD.",
        "Clinical diagnostic performance or superiority over scalar RD.",
        "External replication, corrected full-cohort confirmation, or a claim of being first.",
    ], 7.10, 2.03, 5.18, 3.92, size=16, bullet=True, spacing=14)
    add_footer(slide, 7)

    # 8. Writing sequence
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The figures now dictate the manuscript sequence",
              "Results, Methods and Discussion first; Introduction and Abstract last")
    y = 1.35
    rows = [
        ("Figure 1", "Cohort and design", "Who was analysed, how the connectomes were built, and the three questions"),
        ("Figure 2", "Conventional context", "Show that the cohort reproduces established diffusion abnormalities"),
        ("Figure 3", "Primary stage result", "Report the jointly tested AxD and RD adjacent-stage effects"),
        ("Figure 4", "Systems and sites", "Explain distribution, site deletion and held-out-site RD information"),
        ("Figure 5", "Sensitivity", "State where estimates attenuate and why this limits generalization"),
        ("Supplement", "Alternative dimensions", "Retain graph, edge, selectivity and construct tests without forcing a third story"),
    ]
    for index, (label, heading, body) in enumerate(rows):
        fill = RGBColor(246, 248, 250) if index % 2 == 0 else WHITE
        add_panel(slide, 0.75, y, 11.85, 0.72, fill=fill, border=LINE, radius=False)
        add_text(slide, label, 0.96, y + 0.18, 1.15, 0.28, size=14, color=TEAL, bold=True)
        add_text(slide, heading, 2.22, y + 0.16, 2.35, 0.30, size=16, color=NAVY, bold=True)
        add_text(slide, body, 4.72, y + 0.14, 7.52, 0.36, size=14, color=INK)
        y += 0.83
    add_takeaway(slide, "Next decision with the supervisor: approve this figure order before treating any prose draft as author-ready.")
    add_footer(slide, 8, "Working storyboard | Scientific values copied from validated aggregate outputs")

    prs.core_properties.title = "Structural connectome key results storyboard"
    prs.core_properties.subject = "Figure-led manuscript narrative for supervisor review"
    prs.core_properties.author = "Sabeesh Ethiraj and Dipanjan Roy"
    prs.core_properties.keywords = "ADNI, diffusion MRI, structural connectome, AxD, RD"
    prs.save(PPTX)

    # Fail closed on package integrity and expected content.
    check = Presentation(PPTX)
    if len(check.slides) != 8:
        raise RuntimeError(f"Expected 8 slides, found {len(check.slides)}")
    all_text = "\n".join(
        shape.text
        for slide in check.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    required = [
        "513 participants",
        "AxD β=0.294",
        "RD β=0.591",
        "17/17 AxD",
        "Introduction and Abstract last",
    ]
    missing = [token for token in required if token not in all_text]
    if missing:
        raise RuntimeError(f"Missing required storyboard text: {missing}")
    print(PPTX)


if __name__ == "__main__":
    build()
