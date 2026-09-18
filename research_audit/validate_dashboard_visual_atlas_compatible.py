#!/usr/bin/env python3
"""Independent validation for the Office-compatible dashboard atlas package."""

from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation


ROOT = Path("/home/ec2-user/exp")
OUT = ROOT / "research_audit" / "outputs" / "dashboard_visual_atlas_v1_compatible"
INVENTORY = OUT / "dashboard_visual_inventory_compatible.csv"
VALIDATION = OUT / "validation.json"
OPENXML_PROJECT = Path("/tmp/openxmlcheck/openxmlcheck.csproj")

EXPECTED = {
    "01_structural_connectome_core_review_deck_v1.pptx": 68,
    "02_structural_connectome_image_appendix_A_v1.pptx": 48,
    "03_structural_connectome_image_appendix_B_v1.pptx": 58,
    "structural_connectome_dashboard_visual_atlas_v1_FULL.pptx": 171,
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def inspect_pptx(path: Path, expected_slides: int) -> dict:
    with zipfile.ZipFile(path) as package:
        bad_member = package.testzip()
    presentation = Presentation(path)
    out_of_bounds = []
    pictures = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if int(shape.shape_type) == 13:
                pictures += 1
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > presentation.slide_width + 100
                or shape.top + shape.height > presentation.slide_height + 100
            ):
                out_of_bounds.append({"slide": slide_number, "shape": shape.name})
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "zip_integrity": bad_member is None,
        "cross_engine_reopen": True,
        "slides": len(presentation.slides),
        "expected_slides": expected_slides,
        "slide_count_matches": len(presentation.slides) == expected_slides,
        "picture_count": pictures,
        "all_shapes_within_bounds": not out_of_bounds,
        "out_of_bounds_shapes": out_of_bounds,
    }


def main() -> None:
    deck_paths = [OUT / name for name in EXPECTED]
    deck_results = [inspect_pptx(path, EXPECTED[path.name]) for path in deck_paths]

    command = [
        "dotnet",
        "run",
        "--project",
        str(OPENXML_PROJECT),
        "--configuration",
        "Release",
        "--no-build",
        "--",
        *map(str, deck_paths),
    ]
    strict = subprocess.run(command, capture_output=True, text=True, timeout=180)
    (OUT / "strict_openxml_validation.txt").write_text(strict.stdout + strict.stderr, encoding="utf-8")
    zero_error_records = strict.stdout.count("ERROR_COUNT\t0")

    with INVENTORY.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    core_rows = [row for row in rows if row["deck"] == "01_structural_connectome_core_review_deck_v1.pptx"]
    split_appendix_rows = [
        row
        for row in rows
        if row["deck"]
        in {
            "02_structural_connectome_image_appendix_A_v1.pptx",
            "03_structural_connectome_image_appendix_B_v1.pptx",
        }
        and row["role"] == "appendix_complete"
    ]
    full_appendix_rows = [
        row
        for row in rows
        if row["deck"] == "structural_connectome_dashboard_visual_atlas_v1_FULL.pptx"
        and row["role"] == "appendix_complete"
    ]
    split_sources = {row["source_path"] for row in split_appendix_rows}
    full_sources = {row["source_path"] for row in full_appendix_rows}
    core_live = [row for row in core_rows if row["role"] == "main_live"]

    checks = {
        "all_four_decks_exist": all(path.exists() for path in deck_paths),
        "all_zip_packages_integral": all(item["zip_integrity"] for item in deck_results),
        "all_decks_reopen_with_independent_python_pptx": all(item["cross_engine_reopen"] for item in deck_results),
        "all_slide_counts_match": all(item["slide_count_matches"] for item in deck_results),
        "all_shapes_within_slide_bounds": all(item["all_shapes_within_bounds"] for item in deck_results),
        "microsoft_openxml_sdk_zero_errors_all_decks": strict.returncode == 0 and zero_error_records == len(deck_paths),
        "core_contains_all_54_live_exports": len(core_live) == 54 and len({row["source_path"] for row in core_live}) == 54,
        "split_appendices_cover_245_unique_stored_pngs": len(split_appendix_rows) == 245 and len(split_sources) == 245,
        "full_appendix_covers_245_unique_stored_pngs": len(full_appendix_rows) == 245 and len(full_sources) == 245,
        "split_and_full_appendix_sources_match": split_sources == full_sources,
    }
    record = {
        "record_type": "dashboard_visual_atlas_v1_office_compatible_validation",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "status": "PASS" if all(checks.values()) else "FAIL",
        "generation_engine": "PptxGenJS 4.0.1",
        "strict_validator": "Microsoft Open XML SDK 3.3.0",
        "checks": checks,
        "decks": deck_results,
        "inventory": str(INVENTORY),
        "inventory_sha256": sha256(INVENTORY),
        "strict_validator_output": str(OUT / "strict_openxml_validation.txt"),
    }
    VALIDATION.write_text(json.dumps(record, indent=2), encoding="utf-8")
    print(json.dumps(record, indent=2))
    raise SystemExit(0 if record["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
