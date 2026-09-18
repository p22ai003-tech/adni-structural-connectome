#!/usr/bin/env python3
"""Render the project Markdown reports to readable Word documents.

This intentionally implements the small Markdown subset used by the audit
deliverables (headings, lists, tables, images, code blocks and paragraphs).
The Markdown files remain the source of truth.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


IMAGE_RE = re.compile(r"^!\[([^]]*)\]\(([^)]+)\)\s*$")
LINK_RE = re.compile(r"\[([^]]+)\]\(([^)]+)\)")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    # CT_TcPr has a strict child order.  In particular, w:shd must precede
    # w:noWrap/w:tcMar/w:textDirection/w:tcFitText/w:vAlign/w:hideMark.
    # Appending it after python-docx has emitted w:vAlign produces a package
    # that is well-formed XML but that Microsoft Word reports as unreadable.
    tc_pr.insert_element_before(
        shd,
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
        "w:cellIns",
        "w:cellDel",
        "w:cellMerge",
        "w:tcPrChange",
    )


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    """Add links, bold and inline-code with conservative parsing."""
    cursor = 0
    token = re.compile(r"\[([^]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`")
    for match in token.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        if match.group(1) is not None:
            add_hyperlink(paragraph, match.group(1), match.group(2))
        elif match.group(3) is not None:
            paragraph.add_run(match.group(3)).bold = True
        else:
            run = paragraph.add_run(match.group(4))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def is_table_separator(line: str) -> bool:
    cells = [x.strip() for x in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", x) for x in cells)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if is_table_separator(line):
            continue
        rows.append([x.strip() for x in line.strip().strip("|").split("|")])
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [
        ("Title", 20, "17324D"),
        ("Heading 1", 16, "17324D"),
        ("Heading 2", 13, "2A5D7C"),
        ("Heading 3", 11, "2A5D7C"),
        ("Heading 4", 10, "4C6475"),
    ]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    if "Code Block" not in [s.name for s in document.styles]:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Courier New"
        code.font.size = Pt(8)
        code.paragraph_format.left_indent = Inches(0.25)
        code.paragraph_format.space_after = Pt(3)


def render(source: Path, destination: Path) -> None:
    document = Document()
    configure(document)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    first_heading = True

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue
        if stripped in {"---", "***", "___"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), "B8C4CE")
            border.append(bottom)
            # CT_PPr also has a strict sequence: paragraph borders precede
            # shading, tabs, indentation, justification and spacing.  Keep the
            # element in schema order so desktop Word does not need recovery.
            paragraph._p.get_or_add_pPr().insert_element_before(
                border,
                "w:shd",
                "w:tabs",
                "w:suppressAutoHyphens",
                "w:kinsoku",
                "w:wordWrap",
                "w:overflowPunct",
                "w:topLinePunct",
                "w:autoSpaceDE",
                "w:autoSpaceDN",
                "w:bidi",
                "w:adjustRightInd",
                "w:snapToGrid",
                "w:spacing",
                "w:ind",
                "w:contextualSpacing",
                "w:mirrorIndents",
                "w:suppressOverlap",
                "w:jc",
                "w:textDirection",
                "w:textAlignment",
                "w:textboxTightWrap",
                "w:outlineLvl",
                "w:divId",
                "w:cnfStyle",
                "w:rPr",
                "w:sectPr",
                "w:pPrChange",
            )
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and first_heading:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text)
                first_heading = False
            else:
                paragraph = document.add_heading(level=min(level, 4))
                add_inline(paragraph, text)
            i += 1
            continue

        image_match = IMAGE_RE.match(stripped)
        if image_match:
            image_path = (source.parent / image_match.group(2)).resolve()
            if image_path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.8))
                if image_match.group(1):
                    caption = document.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption.add_run(image_match.group(1))
                    run.italic = True
                    run.font.size = Pt(9)
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, f"[Missing image: {image_match.group(2)}]")
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table(table_lines)
            if rows:
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        cell = table.cell(row_index, col_index)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        paragraph = cell.paragraphs[0]
                        add_inline(paragraph, value)
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                            if row_index == 0:
                                run.bold = True
                        if row_index == 0:
                            set_cell_shading(cell, "DCE6F1")
                document.add_paragraph()
            continue

        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        number = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet or number:
            style = "List Bullet" if bullet else "List Number"
            paragraph = document.add_paragraph(style=style)
            add_inline(paragraph, (bullet or number).group(1))
            i += 1
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            run = paragraph.add_run(stripped.lstrip("> "))
            run.italic = True
            run.font.color.rgb = RGBColor(80, 96, 110)
            i += 1
            continue

        # Join wrapped prose until the next block marker.
        prose = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("```")
                or nxt.startswith("|")
                or nxt.startswith(">")
                or re.match(r"^[-*+]\s+", nxt)
                or re.match(r"^\d+[.)]\s+", nxt)
                or IMAGE_RE.match(nxt)
                or nxt in {"---", "***", "___"}
            ):
                break
            prose.append(nxt)
            i += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(prose))

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(
        f"Generated from {source.name}; Markdown source is authoritative."
    ).font.size = Pt(8)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    render(args.source.resolve(), args.destination.resolve())
    print(args.destination.resolve())


if __name__ == "__main__":
    main()
